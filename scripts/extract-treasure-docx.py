"""Convert approved Treasure DOCX chapters into static website content."""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "treasures"
OUTPUT_PATH = ROOT / "content" / "treasure-data.json"

META = {
    1: ("set-formula-guide", "เซต", "∪", "set"),
    2: ("logic-formula-guide", "ตรรกศาสตร์", "⇒", "logic"),
    3: ("real-numbers-formula-guide", "จำนวนจริงและพหุนาม", "ℝ", "real-numbers"),
    4: ("relations-functions-formula-guide", "ความสัมพันธ์และฟังก์ชัน", "f", "relations-functions"),
    5: ("exponential-logarithm-formula-guide", "เอกซ์โพเนนเชียลและลอการิทึม", "eˣ", "exponential-logarithm"),
    6: ("analytic-geometry-conics-formula-guide", "เรขาคณิตวิเคราะห์และภาคตัดกรวย", "◯", "analytic-geometry"),
    7: ("trigonometry-formula-guide", "ฟังก์ชันตรีโกณมิติ", "θ", "trigonometry"),
    8: ("matrix-formula-guide", "เมทริกซ์", "[ ]", "matrix"),
    9: ("vector-formula-guide", "เวกเตอร์", "→", "vector"),
    10: ("complex-numbers-formula-guide", "จำนวนเชิงซ้อน", "ℂ", "complex-numbers"),
    11: ("counting-probability-formula-guide", "หลักการนับเบื้องต้นและความน่าจะเป็น", "n!", "counting-probability"),
    12: ("sequences-series-formula-guide", "ลำดับและอนุกรม", "Σ", "sequences-series"),
    13: ("calculus-formula-guide", "แคลคูลัสเบื้องต้น", "∫", "calculus"),
    14: ("statistics-distributions-formula-guide", "สถิติและตัวแปรสุ่ม", "σ", "statistics-distributions"),
}

FIELD_LABELS = {
    "ความหมายของตัวแปร:": "meaning",
    "เงื่อนไขการใช้:": "condition",
    "ใช้กับโจทย์แบบไหน:": "use",
    "Keyword:": "keywords",
    "จุดที่ต้องระวัง:": "caution",
    "ข้อผิดพลาดที่ทำบ่อย:": "mistakes",
}


def clean_number(text: str) -> str:
    return re.sub(r"^\d+\.\s*", "", text.strip())


def field_map(text: str) -> dict[str, str]:
    fields: dict[str, str] = {}
    for line in text.splitlines():
        line = line.strip()
        for label, key in FIELD_LABELS.items():
            if line.startswith(label):
                fields[key] = line[len(label):].strip()
                break
    return fields


def table_rows(document: Document, first_cell: str) -> list[list[str]]:
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows and rows[0] and rows[0][0] == first_cell:
            return rows[1:]
    return []


def paragraph_index(paragraphs: list[tuple[str, str]], text: str) -> int:
    return next(i for i, (_, value) in enumerate(paragraphs) if value == text)


def paragraph_index_any(paragraphs: list[tuple[str, str]], *texts: str) -> int:
    return next(i for i, (_, value) in enumerate(paragraphs) if value in texts)


def parse_example(block: list[tuple[str, str]]) -> dict[str, object]:
    values = [value for _, value in block if value]
    header = next((value for value in values if value.startswith("ตัวอย่าง ")), "")
    header_match = re.match(r"ตัวอย่าง\s+\d+\s+(.*?)\s+ระดับ\s+(.+)$", header)
    title = header_match.group(1).strip() if header_match else header
    level = header_match.group(2).strip() if header_match else ""

    def after(label: str) -> str:
        for i, value in enumerate(values):
            if value == label and i + 1 < len(values):
                return values[i + 1]
            if value.startswith(label + ":"):
                return value.split(":", 1)[1].strip()
        return ""

    problem_parts: list[str] = []
    try:
        problem_start = values.index("โจทย์") + 1
        for value in values[problem_start:]:
            if value.startswith("คำตอบ:"):
                break
            problem_parts.append(value)
    except ValueError:
        pass

    steps: list[str] = []
    try:
        start = values.index("วิธีคิดทีละขั้น") + 1
        for value in values[start:]:
            if value.startswith("จุดหลอกของข้อ:"):
                break
            steps.append(clean_number(value))
    except ValueError:
        pass

    return {
        "title": title,
        "level": level,
        "problem": "\n".join(problem_parts) or after("โจทย์"),
        "answer": after("คำตอบ"),
        "formula": after("สูตรที่เลือกใช้"),
        "keywords": after("Keyword สำคัญ"),
        "steps": steps,
        "trap": after("จุดหลอกของข้อ"),
        "check": after("วิธีตรวจคำตอบ"),
    }


def parse_document(path: Path) -> dict[str, object]:
    match = re.search(r"บทที่\s+(\d+)", path.name)
    if not match:
        raise ValueError(f"Chapter number missing: {path.name}")
    chapter = int(match.group(1))
    slug, topic, mark, practice_slug = META[chapter]
    document = Document(path)
    paragraphs = [(p.style.name, p.text.strip()) for p in document.paragraphs]

    overview_i = paragraph_index(paragraphs, "ภาพรวม")
    scope_i = paragraph_index_any(paragraphs, "บทนี้ครอบคลุมอะไร", "บทนี้อยู่ตรงไหนใน A-Level Math 1")
    usage_i = paragraph_index(paragraphs, "วิธีใช้บทนี้ให้คุ้ม")
    formulas_i = paragraph_index(paragraphs, "สูตรสำคัญ")
    patterns_i = paragraph_index(paragraphs, "รูปแบบโจทย์ที่ออกบ่อย")
    traps_i = paragraph_index(paragraphs, "จุดหลอกและข้อควรระวัง")
    scope_note_i = paragraph_index_any(paragraphs, "หมายเหตุด้านขอบเขต", "หัวข้อที่ไม่ควรขยายเกินจำเป็น")
    checklist_i = paragraph_index(paragraphs, "Checklist ก่อนตอบ")
    first_example_i = next(i for i, (_, value) in enumerate(paragraphs) if value == "ตัวอย่างพื้นฐาน")
    references_i = paragraph_index(paragraphs, "แหล่งอ้างอิง")

    overview = [value for _, value in paragraphs[overview_i + 2:scope_i] if value]
    scope = [value for style, value in paragraphs[scope_i + 1:usage_i] if value and style.startswith("List")]

    formulas: list[dict[str, str]] = []
    i = formulas_i + 2
    while i < patterns_i:
        style, value = paragraphs[i]
        if style == "Heading 2" and value:
            expression = paragraphs[i + 1][1]
            details = field_map(paragraphs[i + 2][1])
            formulas.append({"name": clean_number(value), "expression": expression, **details})
            i += 3
        else:
            i += 1

    traps = [clean_number(value) for _, value in paragraphs[traps_i + 2:scope_note_i] if value]
    scope_note = next((value for _, value in paragraphs[scope_note_i + 1:checklist_i] if value), "")

    goal_table = document.tables[0]
    goal = goal_table.cell(0, 0).text.replace("เป้าหมายของบทนี้", "", 1).strip()
    checklist = [
        row[1]
        for table in document.tables
        for row in [[cell.text.strip() for cell in item.cells] for item in table.rows]
        if len(row) > 1 and row[0] == "□"
    ]

    example_starts = [i for i in range(first_example_i, references_i) if paragraphs[i][0] == "Heading 1" and paragraphs[i][1].startswith("ตัวอย่าง")]
    examples = []
    for position, start in enumerate(example_starts):
        end = example_starts[position + 1] if position + 1 < len(example_starts) else references_i
        examples.append(parse_example(paragraphs[start:end]))

    references: list[dict[str, str]] = []
    reference_headings = [
        i for i in range(references_i + 1, len(paragraphs))
        if paragraphs[i][0] == "Heading 2" and re.match(r"^\d+\.", paragraphs[i][1])
    ]
    for position, start in enumerate(reference_headings):
        end = reference_headings[position + 1] if position + 1 < len(reference_headings) else len(paragraphs)
        values = [value for _, value in paragraphs[start + 1:end] if value]
        entry = {"organization": clean_number(paragraphs[start][1])}
        for value in values:
            if value.startswith("เอกสาร:"):
                entry["title"] = value.split(":", 1)[1].strip()
            elif value.startswith("URL:"):
                entry["url"] = value.split(":", 1)[1].strip()
            elif value.startswith("ใช้ตรวจ:"):
                entry["note"] = value.split(":", 1)[1].strip()
        if entry.get("url"):
            references.append(entry)

    patterns = [{"label": row[0], "guidance": row[1]} for row in table_rows(document, "รูปแบบโจทย์") if len(row) > 1]
    keyword_rows = table_rows(document, "คำหรือข้อความในโจทย์") or table_rows(document, "เห็นคำนี้")
    keywords = [{"label": row[0], "guidance": row[1]} for row in keyword_rows if len(row) > 1]
    goal_items = [item.strip() for item in goal.split("•") if item.strip()]
    word_count = sum(len(value.split()) for _, value in paragraphs if value)

    return {
        "chapter": chapter,
        "slug": slug,
        "topic": topic,
        "title": f"คู่มือ {topic}",
        "description": overview[0] if overview else goal,
        "readMinutes": max(10, round(word_count / 150)),
        "mark": mark,
        "practiceSlug": practice_slug,
        "featured": chapter == 1,
        "goal": goal,
        "takeaways": goal_items[:4] or scope[:4],
        "overview": overview,
        "scope": scope,
        "formulas": formulas,
        "patterns": patterns,
        "keywords": keywords,
        "traps": traps,
        "scopeNote": scope_note,
        "checklist": checklist,
        "examples": examples,
        "references": references,
        "sourceDocument": path.name,
    }


def main() -> None:
    documents = sorted(SOURCE_DIR.glob("*.docx"), key=lambda path: int(re.search(r"บทที่\s+(\d+)", path.name).group(1)))
    if len(documents) != 14:
        raise SystemExit(f"Expected 14 chapter documents, found {len(documents)}")
    resources = [parse_document(path) for path in documents]
    OUTPUT_PATH.write_text(json.dumps(resources, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(resources)} chapters to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
